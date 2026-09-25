"""Chat with your files: search inside uploaded documents (PDF, Word, PowerPoint, text, CSV...)."""
from __future__ import annotations

from pathlib import Path

from core.search import bm25, chunk
from core.tools.registry import ToolResult, tool

TEXT_EXT = {".txt", ".md", ".csv", ".json", ".py", ".html", ".xml", ".yaml", ".yml", ".log", ".js", ".sql"}
DOC_EXT = {".pdf", ".docx", ".pptx", ".xlsx"}
_CACHE: dict[tuple[str, int], list[tuple[str, str]]] = {}   # (path, mtime) -> [(location, text)]


def extract(path: Path) -> list[tuple[str, str]]:
    """Return [(location, text)] pieces for a file, e.g. ('page 3', '...')."""
    ext = path.suffix.lower()
    if ext in TEXT_EXT:
        return [("", path.read_text(encoding="utf-8", errors="replace"))]
    if ext == ".pdf":
        from pypdf import PdfReader
        return [(f"page {i}", pg.extract_text() or "") for i, pg in enumerate(PdfReader(str(path)).pages[:300], 1)]
    if ext == ".docx":
        from docx import Document
        doc = Document(str(path))
        paras = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            paras += [" | ".join(c.text for c in row.cells) for row in t.rows]
        return [("", "\n".join(paras))]
    if ext == ".pptx":
        from pptx import Presentation
        out = []
        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
            out.append((f"slide {i}", "\n".join(texts)))
        return out
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            rows = [" | ".join("" if v is None else str(v) for v in row) for row in ws.iter_rows(values_only=True)]
            out.append((f"sheet {ws.title}", "\n".join(rows[:2000])))
        return out
    return []


def index(path: Path) -> list[tuple[str, str]]:
    key = (str(path), path.stat().st_mtime_ns)
    if key not in _CACHE:
        pieces = []
        for loc, text in extract(path):
            pieces += [(loc, c) for c in chunk(text)]
        _CACHE[key] = pieces
    return _CACHE[key]


@tool("search_files",
      "Search inside the user's uploaded files (PDF, Word, PowerPoint, Excel, text, CSV) and return the most "
      "relevant passages with file name and page. Use it to answer questions about the user's documents, "
      "and cite the file and page in your answer.",
      {"query": {"type": "string", "description": "What to look for, in keywords"},
       "filename": {"type": "string", "description": "Optional: search only this file"},
       "top_k": {"type": "integer", "default": 5}},
      ["query"])
def search_files(ctx, query, filename=None, top_k=5):
    files = [p for p in ctx.workspace.files() if p.suffix.lower() in TEXT_EXT | DOC_EXT]
    if filename:
        files = [p for p in files if ctx.workspace.rel(p) == filename or p.name == filename]
    if not files:
        return ToolResult(False, "No searchable files found. Ask the user to upload a document first.")
    passages, errors = [], []
    for p in files:
        try:
            passages += [(ctx.workspace.rel(p), loc, text) for loc, text in index(p)]
        except Exception as e:  # noqa: BLE001
            errors.append(f"{p.name}: {e}")
    ranked = bm25(query, [t for _, _, t in passages])[:max(1, min(int(top_k), 8))]
    if not ranked:
        return ToolResult(True, f"No passages matched '{query}' in {len(files)} file(s). Try other keywords.")
    out = []
    for score, i in ranked:
        name, loc, text = passages[i]
        out.append(f"[{name}{', ' + loc if loc else ''}] (score {score:.1f})\n{text[:900]}")
    note = f"\n\nCould not read: {'; '.join(errors)}" if errors else ""
    return ToolResult(True, f"Top passages from {len(files)} file(s):\n\n" + "\n\n---\n\n".join(out) + note)
