"""Create Excel workbooks (with formulas and charts) and analyze uploaded spreadsheets/CSVs."""
from __future__ import annotations

from core.tools.registry import ToolResult, tool
from core.workspace import safe_filename


def _coerce(v, row=None):
    """Keep formulas as text (filling in {row}); turn numeric-looking text (e.g. '1,250.50') into numbers."""
    if not isinstance(v, str):
        return v
    s = v.strip()
    if s.startswith("="):
        return s.replace("{row}", str(row)) if row else s
    try:
        num = float(s.replace(",", ""))
        return int(num) if num.is_integer() and "." not in s else num
    except ValueError:
        return v


@tool("create_excel",
      "Create an Excel (.xlsx) workbook. Each sheet has column headers (row 1) and data rows (from row 2). "
      "Cells may contain formulas; write {row} for the current row, e.g. '=B{row}*C{row}'. "
      "Set total_row=true to add a correct 'Total' row with SUM formulas for every numeric column. "
      "Optionally add a bar/line/pie chart per sheet (it uses the data rows, not the total).",
      {"filename": {"type": "string"},
       "sheets": {"type": "array", "items": {"type": "object", "properties": {
           "name": {"type": "string"},
           "columns": {"type": "array", "items": {"type": "string"}},
           "rows": {"type": "array", "items": {"type": "array", "items": {}}},
           "total_row": {"type": "boolean", "description": "Add a Total row with SUM formulas"},
           "chart": {"type": "object", "properties": {
               "type": {"type": "string", "enum": ["bar", "line", "pie"]},
               "title": {"type": "string"},
               "label_column": {"type": "integer", "description": "1-based column with labels"},
               "value_columns": {"type": "array", "items": {"type": "integer"},
                                 "description": "1-based columns with numbers"}}}},
           "required": ["name", "columns", "rows"]}}},
      ["filename", "sheets"])
def create_excel(ctx, filename, sheets):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    summary = []
    for spec in sheets:
        ws = wb.create_sheet(str(spec.get("name", "Sheet"))[:31])
        cols = spec.get("columns", [])
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2E6DA4")
            cell.alignment = Alignment(horizontal="center")
        rows = spec.get("rows", [])
        for i, r in enumerate(rows, start=2):
            ws.append([_coerce(v, i) for v in r])
        last_data = ws.max_row
        if spec.get("total_row") and rows:
            total = ["Total"] + [None] * (len(cols) - 1)
            for c in range(2, len(cols) + 1):
                values = [ws.cell(row=i, column=c).value for i in range(2, last_data + 1)]
                if any(isinstance(v, (int, float)) or (isinstance(v, str) and v.startswith("=")) for v in values):
                    letter = get_column_letter(c)
                    total[c - 1] = f"=SUM({letter}2:{letter}{last_data})"
            ws.append(total)
            for c in range(1, len(cols) + 1):
                ws.cell(row=ws.max_row, column=c).font = Font(bold=True)
        for c in range(1, len(cols) + 1):
            width = max(len(str(ws.cell(row=i, column=c).value or "")) for i in range(1, ws.max_row + 1))
            ws.column_dimensions[get_column_letter(c)].width = min(max(10, width + 2), 50)
        ws.freeze_panes = "A2"
        chart_spec = spec.get("chart")
        if chart_spec and ws.max_row > 1:
            kind = chart_spec.get("type", "bar")
            chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}.get(kind, BarChart)()
            chart.title = chart_spec.get("title", "")
            label_col = int(chart_spec.get("label_column", 1))
            value_cols = chart_spec.get("value_columns") or [2]
            for vc in value_cols[:1] if kind == "pie" else value_cols:
                chart.add_data(Reference(ws, min_col=int(vc), min_row=1, max_row=last_data), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=label_col, min_row=2, max_row=last_data))
            chart.height, chart.width = 8, 16
            ws.add_chart(chart, f"{get_column_letter(len(cols) + 2)}2")
        info = f"{ws.title}: data in rows 2-{last_data}, {len(cols)} columns"
        if ws.max_row > last_data:
            info += f", Total in row {ws.max_row}"
        summary.append(info + (f", {kind} chart" if chart_spec else ""))
    path = ctx.workspace.path(safe_filename(filename, ".xlsx"))
    wb.save(path)
    rel = ctx.workspace.rel(path)
    return ToolResult(True, f"Created {rel}: " + "; ".join(summary), files=[rel])


@tool("analyze_file",
      "Load a CSV or Excel file from the workspace (e.g. one the user uploaded) and return its shape, "
      "columns, data types, first rows and summary statistics.",
      {"filename": {"type": "string"},
       "sheet": {"type": "string", "description": "Excel sheet name (optional)"}},
      ["filename"])
def analyze_file(ctx, filename, sheet=None):
    import pandas as pd
    path = ctx.workspace.path(filename)
    if not path.exists():
        return ToolResult(False, f"{filename} not found. Files: {[ctx.workspace.rel(p) for p in ctx.workspace.files()][:20]}")
    if path.suffix.lower() in (".xlsx", ".xlsm", ".xls"):
        df = pd.read_excel(path, sheet_name=sheet or 0)
    else:
        df = pd.read_csv(path)
    info = [f"Shape: {df.shape[0]} rows x {df.shape[1]} columns",
            "Columns and types:\n" + "\n".join(f"  {c}: {t}" for c, t in df.dtypes.astype(str).items()),
            "First rows:\n" + df.head(8).to_string(max_colwidth=40),
            "Summary:\n" + df.describe(include="all").transpose().head(30).to_string(max_colwidth=30)]
    return ToolResult(True, "\n\n".join(info))
