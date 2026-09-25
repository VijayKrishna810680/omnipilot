"""Create Word, PDF and PowerPoint files from simple Markdown written by the AI."""
from __future__ import annotations

import re

from core.tools.registry import ToolResult, tool
from core.workspace import safe_filename


# ---------------------------------------------------------------- Markdown parsing
def parse_markdown(md: str) -> list[tuple]:
    """Turn a small Markdown subset into blocks: heading, bullet, number, table, para."""
    blocks, table = [], []
    for raw in md.splitlines():
        line = raw.rstrip()
        if line.strip().startswith("|") and line.strip().endswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                table.append(cells)
            continue
        if table:
            blocks.append(("table", table))
            table = []
        if not line.strip():
            continue
        if m := re.match(r"^(#{1,4})\s+(.*)", line):
            blocks.append(("heading", len(m.group(1)), m.group(2).strip()))
        elif m := re.match(r"^\s*[-*•]\s+(.*)", line):
            blocks.append(("bullet", m.group(1).strip()))
        elif m := re.match(r"^\s*\d+[.)]\s+(.*)", line):
            blocks.append(("number", m.group(1).strip()))
        else:
            blocks.append(("para", line.strip()))
    if table:
        blocks.append(("table", table))
    return blocks


def _runs(text: str):
    """Split '**bold**' segments: yields (text, is_bold)."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            yield part, i % 2 == 1


# ---------------------------------------------------------------- Word
def _write_docx(path, title, blocks):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    if title:
        doc.add_heading(title, level=0)
    for b in blocks:
        kind = b[0]
        if kind == "heading":
            doc.add_heading(b[2], level=min(b[1], 4))
        elif kind in ("bullet", "number", "para"):
            style = {"bullet": "List Bullet", "number": "List Number", "para": None}[kind]
            p = doc.add_paragraph(style=style)
            for text, bold in _runs(b[1]):
                p.add_run(text).bold = bold
        elif kind == "table":
            rows = b[1]
            width = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=width)
            t.style = "Light Grid Accent 1"
            for i, r in enumerate(rows):
                for j in range(width):
                    t.cell(i, j).text = r[j] if j < len(r) else ""
                    if i == 0:
                        for run in t.cell(i, j).paragraphs[0].runs:
                            run.bold = True
    doc.save(path)


# ---------------------------------------------------------------- PDF
def _write_pdf(path, title, blocks):
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()

    def fmt(text):
        return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escape(text))

    story = []
    if title:
        story += [Paragraph(fmt(title), styles["Title"]), Spacer(1, 0.3 * cm)]
    pending_list, list_kind = [], None

    def flush():
        nonlocal pending_list, list_kind
        if pending_list:
            story.append(ListFlowable(pending_list, bulletType="1" if list_kind == "number" else "bullet"))
            pending_list, list_kind = [], None

    for b in blocks:
        kind = b[0]
        if kind in ("bullet", "number"):
            if list_kind not in (None, kind):
                flush()
            list_kind = kind
            pending_list.append(ListItem(Paragraph(fmt(b[1]), styles["BodyText"])))
            continue
        flush()
        if kind == "heading":
            story.append(Paragraph(fmt(b[2]), styles[{1: "Heading1", 2: "Heading2"}.get(b[1], "Heading3")]))
        elif kind == "para":
            story.append(Paragraph(fmt(b[1]), styles["BodyText"]))
        elif kind == "table":
            rows = b[1]
            width = max(len(r) for r in rows)
            data = [[Paragraph(fmt(c), styles["BodyText"]) for c in (r + [""] * (width - len(r)))] for r in rows]
            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dbe7f5")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story += [t, Spacer(1, 0.3 * cm)]
    flush()
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                      topMargin=2 * cm, bottomMargin=2 * cm, title=title or "").build(story)


@tool("create_document",
      "Create a Word (.docx) or PDF document from Markdown text. Supports # headings, - bullets, "
      "1. numbered lists, **bold** and | tables |. Use for reports, letters, resumes, notes, plans.",
      {"format": {"type": "string", "enum": ["docx", "pdf"]},
       "filename": {"type": "string", "description": "File name, e.g. 'sales_report'"},
       "title": {"type": "string"},
       "content_markdown": {"type": "string", "description": "The full document body in Markdown"}},
      ["format", "filename", "content_markdown"])
def create_document(ctx, format, filename, content_markdown, title=""):
    fmt = format.lower().strip(".")
    if fmt not in ("docx", "pdf"):
        return ToolResult(False, "format must be 'docx' or 'pdf'")
    path = ctx.workspace.path(safe_filename(filename, "." + fmt))
    blocks = parse_markdown(content_markdown)
    (_write_docx if fmt == "docx" else _write_pdf)(path, title, blocks)
    rel = ctx.workspace.rel(path)
    return ToolResult(True, f"Created {rel} ({len(blocks)} blocks, {path.stat().st_size // 1024 + 1} KB)", files=[rel])


@tool("create_presentation",
      "Create a PowerPoint (.pptx) slide deck. Each slide has a title, bullet points and optional speaker notes.",
      {"filename": {"type": "string"},
       "title": {"type": "string", "description": "Title of the first slide"},
       "subtitle": {"type": "string"},
       "slides": {"type": "array", "items": {"type": "object", "properties": {
           "title": {"type": "string"},
           "bullets": {"type": "array", "items": {"type": "string"}},
           "notes": {"type": "string"}}, "required": ["title"]}}},
      ["filename", "title", "slides"])
def create_presentation(ctx, filename, title, slides, subtitle=""):
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[0])
    first.shapes.title.text = title
    if len(first.placeholders) > 1:
        first.placeholders[1].text = subtitle
    for s in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = s.get("title", "")
        body = slide.placeholders[1].text_frame
        bullets = s.get("bullets") or []
        for i, b in enumerate(bullets):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = str(b)
            para.font.size = Pt(20 if len(bullets) <= 5 else 16)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
    path = ctx.workspace.path(safe_filename(filename, ".pptx"))
    prs.save(path)
    rel = ctx.workspace.rel(path)
    return ToolResult(True, f"Created {rel} with {len(slides) + 1} slides", files=[rel])
