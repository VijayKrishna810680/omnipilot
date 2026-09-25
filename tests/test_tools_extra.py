"""Excel totals, context-window trimming and running tests inside the sandbox."""
import json

from openpyxl import load_workbook

from core.agent import Agent, Context
from core.tools.registry import run_tool
from tests.conftest import ScriptedRouter


def test_excel_total_row_and_row_formulas(env):
    ws, mem, _ = env
    args = {"filename": "budget", "sheets": [{
        "name": "Budget", "columns": ["Item", "Planned", "Actual", "Diff"],
        "rows": [["Rent", 1500, 1400, "=B{row}-C{row}"], ["Food", "600", 650, "=B{row}-C{row}"],
                 ["Other", 200, 100, "=B{row}-C{row}"]],
        "total_row": True, "chart": {"type": "pie", "label_column": 1, "value_columns": [2]}}]}
    res = run_tool("create_excel", json.dumps(args), Context(ws, mem))
    assert res.ok and "Total in row 5" in res.output
    assert "Planned=2,300" in res.output and "Actual=2,150" in res.output
    sheet = load_workbook(ws.path("budget.xlsx"))["Budget"]
    assert sheet["D3"].value == "=B3-C3"
    assert [sheet.cell(row=5, column=c).value for c in range(1, 5)] == ["Total", "=SUM(B2:B4)", "=SUM(C2:C4)", "=SUM(D2:D4)"]


def test_window_trims_old_file_contents_and_stays_valid(env):
    ws, mem, log = env
    agent = Agent(ScriptedRouter([]), ws, mem, log=log)
    big = "x = 1\n" * 3000
    for i in range(4):
        agent.messages += [
            {"role": "user", "content": f"task {i}"},
            {"role": "assistant", "content": "", "tool_calls": [{"id": f"c{i}", "type": "function", "function": {
                "name": "write_file", "arguments": json.dumps({"path": f"a{i}.py", "content": big})}}]},
            {"role": "tool", "tool_call_id": f"c{i}", "content": "y" * 5000},
            {"role": "assistant", "content": "done"}]
    window = agent._window()
    assert window[1]["role"] == "user"
    assert sum(len(json.dumps(m)) for m in window[1:]) <= 14000
    for m in window:
        for c in m.get("tool_calls", []):
            json.loads(c["function"]["arguments"])  # still valid JSON
    assert window[-1]["content"] == "done"


def test_pytest_runs_inside_sandbox(tmp_path):
    from core.sandbox import run_python
    (tmp_path / "app").mkdir()
    (tmp_path / "app" / "todo.py").write_text("def add(xs, x):\n    return xs + [x]\n")
    (tmp_path / "app" / "test_todo.py").write_text("from todo import add\n\ndef test_add():\n    assert add([], 1) == [1]\n")
    r = run_python("import pytest; raise SystemExit(pytest.main(['-q', '-p', 'no:cacheprovider', 'app']))", tmp_path)
    assert r.ok, r.stdout + r.stderr
    assert "1 passed" in r.stdout
    assert not any("__pycache__" in f for f in r.new_files)
