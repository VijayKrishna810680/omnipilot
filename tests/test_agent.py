"""End-to-end agent tests with a scripted model: every tool, approvals, memory, errors."""
import zipfile

from conftest import ScriptedRouter

from core.agent import Agent


def run(agent, text):
    return list(agent.send(text))


def test_documents_excel_slides_in_one_task(env):
    ws, mem, log = env
    router = ScriptedRouter([
        [("create_document", {"format": "docx", "filename": "report", "title": "Sales Report",
                              "content_markdown": "# Summary\nSales grew **12%**.\n- North: 40\n- South: 60\n\n| Region | Sales |\n|---|---|\n| North | 40 |\n| South | 60 |"}),
         ("create_document", {"format": "pdf", "filename": "report", "title": "Sales Report",
                              "content_markdown": "# Summary\n1. First\n2. Second\n\n| A | B |\n|---|---|\n| 1 | 2 |"}),
         ("create_excel", {"filename": "sales", "sheets": [{"name": "Data", "columns": ["Region", "Sales"],
                           "rows": [["North", "40"], ["South", "60"], ["Total", "=SUM(B2:B3)"]],
                           "chart": {"type": "bar", "title": "Sales", "label_column": 1, "value_columns": [2]}}]}),
         ("create_presentation", {"filename": "deck", "title": "Q3 Review", "slides": [
             {"title": "Highlights", "bullets": ["Sales +12%", "2 regions"], "notes": "Keep short"}]})],
        "Done: report.docx, report.pdf, sales.xlsx and deck.pptx.",
    ])
    events = run(Agent(router, ws, mem, log=log), "Make a sales report, excel and slides")
    results = [e for e in events if e.type == "tool_result"]
    assert all(r.data["ok"] for r in results), [r.data for r in results]
    assert events[-1].type == "final"

    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation
    d = Document(ws.path("report.docx"))
    assert any("12%" in p.text for p in d.paragraphs) and len(d.tables) == 1
    assert ws.path("report.pdf").read_bytes()[:4] == b"%PDF"
    wb = load_workbook(ws.path("sales.xlsx"))
    sheet = wb["Data"]
    assert sheet["B2"].value == 40 and sheet["B4"].value == "=SUM(B2:B3)" and len(sheet._charts) == 1
    assert len(Presentation(ws.path("deck.pptx")).slides) == 2


def test_code_needs_approval_then_resumes(env):
    ws, mem, log = env
    router = ScriptedRouter([
        [("write_file", {"path": "calc/app.py", "content": "def add(a, b):\n    return a + b\n"}),
         ("run_python", {"code": "import sys; sys.path.insert(0, 'calc')\nfrom app import add\nprint(add(2, 3))"})],
        [("zip_project", {"folder": "calc"})],
        "Project built, tested (2+3=5) and zipped as calc.zip.",
    ])
    agent = Agent(router, ws, mem, log=log)
    events = run(agent, "Build a calculator project")
    assert events[-1].type == "approval" and events[-1].data["tool"] == "run_python"
    assert ws.path("calc/app.py").exists()          # safe tool ran before the pause
    events = list(agent.resume(approved=True))
    outputs = [e.data for e in events if e.type == "tool_result"]
    assert outputs[0]["output"].strip() == "5"
    assert zipfile.ZipFile(ws.path("calc.zip")).namelist() == ["calc/app.py"]
    assert events[-1].type == "final"


def test_denied_code_is_not_run(env):
    ws, mem, log = env
    router = ScriptedRouter([[("run_python", {"code": "open('pwned.txt','w').write('x')"})],
                             "Okay, I will not run it."])
    agent = Agent(router, ws, mem, log=log)
    run(agent, "run something")
    events = list(agent.resume(approved=False))
    assert not ws.path("pwned.txt").exists()
    assert "denied" in router.seen[-1][-1]["content"]
    assert events[-1].type == "final"


def test_chart_from_sandbox_is_shown(env):
    ws, mem, log = env
    router = ScriptedRouter([[("run_python", {"code": "import matplotlib.pyplot as plt\nplt.bar(['a','b'],[3,5])\nplt.savefig('chart.png')"})],
                             "Here is your chart."])
    agent = Agent(router, ws, mem, log=log, auto_approve=True)
    result = [e for e in run(agent, "chart") if e.type == "tool_result"][0]
    assert result.data["ok"] and result.data["image"] == "chart.png"


def test_memory_is_saved_and_injected(env):
    ws, mem, log = env
    router = ScriptedRouter([[("remember", {"fact": "User prefers reports in PDF format"})],
                             "Noted!", "Sure, PDF it is."])
    agent = Agent(router, ws, mem, log=log)
    run(agent, "I always want PDF reports")
    run(agent, "make me a report")
    assert "PDF" in router.seen[-1][0]["content"]      # system prompt contains the memory
    assert mem.search("report format") == ["User prefers reports in PDF format"]


def test_bad_tool_input_becomes_error_not_crash(env):
    ws, mem, log = env
    router = ScriptedRouter([[("create_document", {"format": "exe", "filename": "x", "content_markdown": "hi"}),
                              ("no_such_tool", {}),
                              ("write_file", {"path": "../../etc/passwd", "content": "x"})],
                             "Sorry, those failed."])
    results = [e.data for e in run(Agent(router, ws, mem, log=log), "x") if e.type == "tool_result"]
    assert [r["ok"] for r in results] == [False, False, False]
    assert "outside the workspace" in results[2]["output"]


def test_no_provider_gives_clear_error(env, monkeypatch):
    from core.router import Router
    for k in ("GROQ_API_KEY", "GEMINI_API_KEY", "OPENROUTER_API_KEY", "OPENAI_API_KEY"):
        monkeypatch.delenv(k, raising=False)
    ws, mem, log = env
    events = run(Agent(Router(enable_ollama=False), ws, mem, log=log), "hi")
    assert events[-1].type == "error" and "No AI provider" in events[-1].data["message"]


def test_long_chat_window_starts_with_user(env):
    ws, mem, log = env
    agent = Agent(ScriptedRouter([]), ws, mem, log=log)
    for i in range(60):
        agent.messages += [{"role": "user", "content": f"q{i}"}, {"role": "assistant", "content": f"a{i}"}]
    window = agent._window(max_messages=15)
    assert window[0]["role"] == "system" and window[1]["role"] == "user" and len(window) <= 16
